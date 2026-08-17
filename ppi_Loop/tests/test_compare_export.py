"""测试单一多模型协作结果的标准 Excel 导出。"""

import pandas as pd
from openpyxl import load_workbook

from ppi_loop.compare_export import (
    STANDARD_OUTPUT_COLS,
    build_standard_output,
    export_standard_output,
)


def _raw_data() -> pd.DataFrame:
    return pd.DataFrame([
        {
            'row_id': '1', 'bus_date': '2026-08-11',
            'brand_name_product': 'SHISEIDO', 'item_id': 'A1',
            'product_title': '标题1', 'image_url': 'https://example.com/1.jpg',
            'label': '红腰子精华', 'type': '单品',
            'RSP': '999', 'RSP with PWP': '888',
            'Actual Price': '777', 'PPI': '0.5',
            'Unnamed: 12': None,
        },
        {
            'row_id': '2', 'bus_date': '2026-08-11',
            'brand_name_product': 'SHISEIDO', 'item_id': 'A2',
            'product_title': '标题2', 'image_url': 'https://example.com/2.jpg',
            'label': '悦薇水', 'type': '单品',
        },
    ])


def _model_output() -> pd.DataFrame:
    return pd.DataFrame([
        {
            '主图URL': 'https://example.com/1.jpg', '组合ID': 'S1',
            '类型': '主品', '名称': '红腰子精华', '规格': '50', '单位': 'ml',
            '数量': '1', '金额': '', '使用条件': '',
            'RapidFuzz匹配分数': 100.0, 'VLM测评分数': 96,
        },
        {
            '主图URL': 'https://example.com/1.jpg', '组合ID': 'S1',
            '类型': '赠品', '名称': '红腰子精华', '规格': '10', '单位': 'ml',
            '数量': '2', '金额': '', '使用条件': '',
            'RapidFuzz匹配分数': 98.0, 'VLM测评分数': 92,
        },
    ])


def test_build_standard_output_keeps_quantity_and_omits_unit():
    result = build_standard_output(_raw_data(), _model_output())

    assert list(result.columns) == STANDARD_OUTPUT_COLS
    assert len(result) == 3
    assert '数量' in result.columns
    assert '单位' not in result.columns
    assert '主图URL' not in result.columns
    assert list(result.iloc[:2]['数量']) == ['1', '2']
    assert list(result.iloc[:2]['RapidFuzz匹配分数']) == [100.0, 98.0]
    assert list(result.iloc[:2]['VLM测评分数']) == [96.0, 92.0]
    assert result.iloc[2]['image_url'] == 'https://example.com/2.jpg'
    assert pd.isna(result.iloc[2]['类型'])
    assert result[['RSP', 'RSP with PWP', 'Actual Price', 'PPI']].isna().all().all()


def test_export_standard_output_writes_grouped_headers(tmp_path):
    output_path = tmp_path / 'standard.xlsx'
    export_standard_output(_raw_data(), _model_output(), output_path)

    workbook = load_workbook(output_path, data_only=False)
    sheet = workbook['Sheet1']

    assert sheet.max_column == 37
    assert sheet.freeze_panes == 'A3'
    assert sheet['M1'].value == 'PPI_LANGGRAPH'
    assert sheet['T1'].value == 'PPI_LANGGRAPH - 结果检验'
    assert sheet['M2'].value == '组合ID'
    assert sheet['Q2'].value == '数量'
    assert sheet['AF2'].value == 'RapidFuzz匹配分数'
    assert sheet['AG2'].value == 'VLM测评分数'
    assert 'A1:A2' in {str(r) for r in sheet.merged_cells.ranges}
    assert 'M1:S1' in {str(r) for r in sheet.merged_cells.ranges}
    assert 'T1:AG1' in {str(r) for r in sheet.merged_cells.ranges}
    assert 'AH1:AH2' in {str(r) for r in sheet.merged_cells.ranges}
    assert sheet['A3'].value == '1'
    assert sheet['M3'].value == 'S1'
    assert sheet['Q3'].value == '1'


def test_empty_model_output_still_keeps_source_rows():
    result = build_standard_output(_raw_data(), pd.DataFrame())

    assert len(result) == 2
    assert result['image_url'].tolist() == [
        'https://example.com/1.jpg',
        'https://example.com/2.jpg',
    ]
    assert result[['组合ID', '类型', '名称']].isna().all().all()
